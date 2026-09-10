from flask import Flask, jsonify, request, send_from_directory
from pathlib import Path
from datetime import datetime, timezone
import json, ssl, threading, time, uuid
from openpyxl import load_workbook
import paho.mqtt.client as mqtt
import config

BASE = Path(__file__).resolve().parent
DB_PATH = BASE / config.DATABASE_FILE
app = Flask(__name__, static_folder="static", static_url_path="/static")
lock = threading.RLock()

state = {
    "brokerConnected": False,
    "lastBrokerEvent": None,
    "locker": {"id":"A","status":"offline","door":"unknown","battery":None,"online":False,"lastSeen":None,"currentBooking":""},
    "lastAlert": None,
}

def now_iso():
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")

def open_wb(**kwargs):
    return load_workbook(DB_PATH, **kwargs)

def append_row(sheet_name, row):
    with lock:
        wb = open_wb(); ws = wb[sheet_name]; ws.append(row); wb.save(DB_PATH)

def update_locker(**updates):
    with lock:
        wb = open_wb(); ws = wb["Lockers"]
        row = None
        for r in range(2, ws.max_row+1):
            if str(ws.cell(r,1).value) == "A": row = r; break
        if row is None:
            row = ws.max_row+1; ws.cell(row,1).value="A"; ws.cell(row,2).value="Beach Zone 1"
        mapping = {"status":3,"door":4,"battery":5,"online":6,"lastSeen":7,"hourlyRate":8,"currentBooking":9}
        for k,v in updates.items():
            if k in mapping: ws.cell(row,mapping[k]).value = v
        wb.save(DB_PATH)
    state["locker"].update(updates)

def upsert_user(data):
    mobile = str(data.get("mobile") or "")
    if not mobile: return
    with lock:
        wb = open_wb(); ws=wb["Users"]; row=None
        for r in range(2, ws.max_row+1):
            if str(ws.cell(r,3).value or "") == mobile: row=r; break
        ts=now_iso()
        if row is None:
            row=ws.max_row+1
            ws.cell(row,1).value = data.get("userId") or f"USR-{mobile[-6:]}"
            ws.cell(row,5).value = ts
        ws.cell(row,2).value = data.get("name") or ws.cell(row,2).value
        ws.cell(row,3).value = mobile
        ws.cell(row,4).value = data.get("paymentLast4") or ws.cell(row,4).value
        ws.cell(row,6).value = ts
        ws.cell(row,7).value = data.get("status") or "active"
        wb.save(DB_PATH)

def upsert_reservation(data):
    bid = str(data.get("bookingId") or "")
    if not bid: return
    with lock:
        wb=open_wb(); ws=wb["Reservations"]; row=None
        for r in range(2,ws.max_row+1):
            if str(ws.cell(r,1).value or "") == bid: row=r; break
        if row is None: row=ws.max_row+1
        vals = [
            bid, data.get("locker","A"), data.get("ownerMobile",""), data.get("ownerName",""), data.get("date") or data.get("localDate",""),
            data.get("from") or data.get("localStart",""), data.get("to") or data.get("localEnd",""), data.get("startIso",""), data.get("endIso",""),
            data.get("pin",""), data.get("hourlyRate",0), data.get("total") or data.get("paidAmount",0), data.get("lateFee",0), data.get("finalTotal",0),
            data.get("status","reserved"), data.get("paymentStatus","paid-demo"), data.get("createdAt") or data.get("issuedAt") or now_iso(), data.get("completedAt",""), data.get("cancelledAt","")
        ]
        for c,v in enumerate(vals,1): ws.cell(row,c).value=v
        wb.save(DB_PATH)

def add_access(data):
    append_row("Access Logs", [data.get("timestamp") or now_iso(), data.get("locker","A"), data.get("bookingId",""), data.get("ownerMobile",""), data.get("method",""), data.get("action",""), data.get("result",""), data.get("door",""), data.get("source","app")])

def log_device(topic, payload):
    suffix = topic.split(f"{config.BASE_TOPIC}/",1)[-1] if f"{config.BASE_TOPIC}/" in topic else topic
    value = payload
    if isinstance(payload, dict): value = payload.get("value", payload.get("state", payload.get("status", payload.get("door", payload.get("battery", "")))))
    append_row("Device Events", [now_iso(), "A", topic, suffix, str(value), json.dumps(payload, ensure_ascii=False) if not isinstance(payload,str) else payload])

def parse_payload(msg):
    text=msg.payload.decode("utf-8",errors="replace")
    if text=="": return ""
    try: return json.loads(text)
    except: return text

def on_connect(client, userdata, flags, reason_code, properties=None):
    state["brokerConnected"] = True; state["lastBrokerEvent"] = now_iso()
    client.subscribe(f"{config.BASE_TOPIC}/#", qos=config.QOS)
    update_locker(online=True,lastSeen=now_iso())

def on_disconnect(client, userdata, flags, reason_code, properties=None):
    state["brokerConnected"] = False; state["lastBrokerEvent"] = now_iso(); update_locker(online=False)

def on_message(client, userdata, msg):
    payload = parse_payload(msg); topic=msg.topic; state["lastBrokerEvent"] = now_iso()
    log_device(topic,payload)
    suffix = topic.split(f"{config.BASE_TOPIC}/",1)[-1]
    if suffix == "status":
        val = payload.get("status",payload.get("state",payload.get("value","online"))) if isinstance(payload,dict) else payload
        update_locker(status=str(val).lower(), online=True,lastSeen=now_iso())
    elif suffix == "door":
        val = payload.get("door",payload.get("state",payload.get("value","unknown"))) if isinstance(payload,dict) else payload
        update_locker(door=str(val).lower(), online=True,lastSeen=now_iso())
    elif suffix == "battery":
        val = payload.get("battery",payload.get("value",payload.get("percent"))) if isinstance(payload,dict) else payload
        try: val=float(val)
        except: pass
        update_locker(battery=val, online=True,lastSeen=now_iso())
    elif suffix == "alert":
        if isinstance(payload,dict):
            typ=payload.get("type","Device Alert"); sev=payload.get("severity","warning"); message=payload.get("message",json.dumps(payload))
        else: typ="Device Alert"; sev="warning"; message=str(payload)
        append_row("Alerts",[now_iso(),"A",typ,sev,message,False]); state["lastAlert"]={"type":typ,"severity":sev,"message":message,"timestamp":now_iso()}
    elif suffix == "app/user/upsert" and isinstance(payload,dict): upsert_user(payload)
    elif suffix == "app/reservation/upsert" and isinstance(payload,dict):
        upsert_reservation(payload); update_locker(currentBooking=payload.get("bookingId", state["locker"].get("currentBooking","")))
    elif suffix == "app/reservation/cancel" and isinstance(payload,dict): upsert_reservation(payload); update_locker(currentBooking="")
    elif suffix == "app/reservation/complete" and isinstance(payload,dict):
        upsert_reservation(payload); update_locker(currentBooking="")
        if payload.get("lateFee",0): append_row("Payments",[now_iso(),payload.get("bookingId",""),payload.get("ownerMobile",""),"Late Fee",payload.get("lateFee",0),"due-demo",""])
    elif suffix == "app/access" and isinstance(payload,dict): add_access(payload)
    elif suffix == "app/payment" and isinstance(payload,dict): append_row("Payments",[payload.get("timestamp") or now_iso(),payload.get("bookingId",""),payload.get("ownerMobile",""),payload.get("type","Payment"),payload.get("amount",0),payload.get("status","paid-demo"),payload.get("reference","")])
    elif suffix == "app/feedback" and isinstance(payload,dict): append_row("Feedback",[payload.get("timestamp") or now_iso(),payload.get("bookingId",""),payload.get("ownerMobile",""),payload.get("rating",0),payload.get("comment","")])


def find_reservation(booking_id):
    with lock:
        wb=open_wb(read_only=True,data_only=True); ws=wb["Reservations"]
        headers=[c.value for c in ws[1]]
        found=None
        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[0] or "") == str(booking_id):
                found={str(headers[i]): row[i] for i in range(min(len(headers),len(row)))}
                break
        wb.close(); return found

def update_reservation_status(booking_id, status, cancelled_at=None):
    with lock:
        wb=open_wb(); ws=wb["Reservations"]; row=None
        for r in range(2,ws.max_row+1):
            if str(ws.cell(r,1).value or "") == str(booking_id): row=r; break
        if row is None:
            wb.close(); return False
        ws.cell(row,15).value=status
        if cancelled_at is not None: ws.cell(row,19).value=cancelled_at
        wb.save(DB_PATH); return True

def add_admin_audit(action, booking_id="", locker="A", result="ok", details="", admin="Owner Dashboard"):
    append_row("Admin Audit", [now_iso(),admin,action,booking_id,locker,result,details])

def safe_publish(topic, payload):
    try:
        info=mqttc.publish(topic,json.dumps(payload),qos=config.QOS)
        info.wait_for_publish(timeout=3)
        return True
    except Exception:
        return False

mqttc = mqtt.Client(mqtt.CallbackAPIVersion.VERSION2, client_id=f"sandlock-dashboard-{uuid.uuid4().hex[:8]}")
mqttc.username_pw_set(config.BROKER_USERNAME, config.BROKER_PASSWORD)
mqttc.tls_set(tls_version=ssl.PROTOCOL_TLS_CLIENT)
mqttc.on_connect=on_connect; mqttc.on_disconnect=on_disconnect; mqttc.on_message=on_message

def mqtt_loop():
    while True:
        try:
            mqttc.connect(config.BROKER_HOST, config.BROKER_PORT, keepalive=45)
            mqttc.loop_forever(retry_first_connection=True)
        except Exception:
            state["brokerConnected"]=False; time.sleep(3)
threading.Thread(target=mqtt_loop, daemon=True).start()

def rows_as_dict(sheet_name, limit=100):
    with lock:
        wb=open_wb(read_only=True,data_only=True); ws=wb[sheet_name]; headers=[c.value for c in ws[1]]; out=[]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(v not in (None,"") for v in row): continue
            out.append({str(headers[i]):row[i] for i in range(min(len(headers),len(row)))})
        wb.close(); return list(reversed(out[-limit:]))

@app.route("/")
def home(): return send_from_directory(BASE/"static","index.html")
@app.get("/api/overview")
def api_overview():
    reservations=rows_as_dict("Reservations",200); users=rows_as_dict("Users",200); alerts=rows_as_dict("Alerts",50); payments=rows_as_dict("Payments",200)
    active=[r for r in reservations if str(r.get("Status","")).lower() not in ("completed","cancelled","")]
    revenue=sum(float(p.get("Amount EGP") or 0) for p in payments if str(p.get("Status","")).lower() not in ("failed","cancelled"))
    return jsonify({"state":state,"kpi":{"users":len(users),"activeReservations":len(active),"alerts":len([a for a in alerts if not a.get("Acknowledged")]),"revenue":round(revenue,2)},"activeReservations":active[:8],"alerts":alerts[:8]})
@app.get("/api/<name>")
def api_sheet(name):
    mapping={"users":"Users","reservations":"Reservations","access":"Access Logs","events":"Device Events","alerts":"Alerts","payments":"Payments","feedback":"Feedback","lockers":"Lockers","audit":"Admin Audit"}
    if name not in mapping: return jsonify({"error":"not found"}),404
    return jsonify(rows_as_dict(mapping[name],300))
@app.get("/api/reservations/<booking_id>")
def reservation_details(booking_id):
    r=find_reservation(booking_id)
    if not r: return jsonify({"error":"reservation not found"}),404
    return jsonify(r)

@app.post("/api/reservations/<booking_id>/reveal-pin")
def reveal_pin(booking_id):
    r=find_reservation(booking_id)
    if not r: return jsonify({"error":"reservation not found"}),404
    add_admin_audit("Reveal Reservation PIN", booking_id, str(r.get("Locker") or "A"), "success", "PIN revealed from reservation details")
    add_access({"timestamp":now_iso(),"locker":r.get("Locker") or "A","bookingId":booking_id,"ownerMobile":r.get("User Mobile") or "","method":"Admin Dashboard","action":"reveal-pin","result":"success","door":state["locker"].get("door",""),"source":"admin"})
    return jsonify({"ok":True,"pin":str(r.get("PIN") or "")})

@app.post("/api/reservations/<booking_id>/cancel")
def cancel_reservation(booking_id):
    r=find_reservation(booking_id)
    if not r: return jsonify({"error":"reservation not found"}),404
    old=str(r.get("Status") or "").lower()
    if old in ("cancelled","completed"):
        return jsonify({"error":f"reservation is already {old}"}),409
    ts=now_iso()
    if not update_reservation_status(booking_id,"Cancelled",ts): return jsonify({"error":"unable to update reservation"}),500
    locker=str(r.get("Locker") or "A")
    if locker == "A": update_locker(status="available",currentBooking="")
    payload={"bookingId":booking_id,"locker":locker,"status":"Cancelled","cancelledAt":ts,"source":"admin-dashboard"}
    published=safe_publish(f"{config.BASE_TOPIC}/admin/reservation/cancel",payload)
    safe_publish(f"{config.BASE_TOPIC}/cmd/cancel",payload)
    add_admin_audit("Cancel Reservation",booking_id,locker,"success" if published else "database-updated / mqtt-not-confirmed",f"Previous status: {r.get('Status','')}")
    add_access({"timestamp":ts,"locker":locker,"bookingId":booking_id,"ownerMobile":r.get("User Mobile") or "","method":"Admin Dashboard","action":"cancel-reservation","result":"success","door":state["locker"].get("door",""),"source":"admin"})
    return jsonify({"ok":True,"mqttPublished":published})

@app.post("/api/cmd/unlock")
def cmd_unlock():
    payload=request.get_json(silent=True) or {}
    booking_id=str(payload.get("bookingId") or "")
    r=find_reservation(booking_id) if booking_id else None
    locker=str((r or {}).get("Locker") or payload.get("locker") or "A")
    payload.update({"locker":locker,"source":"admin-dashboard","timestamp":now_iso(),"ownerOverride":True})
    published=safe_publish(f"{config.BASE_TOPIC}/cmd/unlock",payload)
    add_access({"timestamp":now_iso(),"locker":locker,"bookingId":booking_id,"ownerMobile":((r or {}).get("User Mobile") or ""),"method":"Admin Dashboard","action":"emergency-unlock","result":"command-sent" if published else "mqtt-not-confirmed","door":state["locker"].get("door",""),"source":"admin"})
    add_admin_audit("Emergency Remote Unlock",booking_id,locker,"command-sent" if published else "mqtt-not-confirmed","Owner override unlock")
    return jsonify({"ok":True,"mqttPublished":published})

@app.post("/api/alerts/ack")
def ack_alert():
    idx=int((request.get_json(silent=True) or {}).get("row",0))
    with lock:
        wb=open_wb(); ws=wb["Alerts"]
        if idx>=2 and idx<=ws.max_row: ws.cell(idx,6).value=True; wb.save(DB_PATH)
    return jsonify({"ok":True})
@app.get("/download/database.xlsx")
def download_db(): return send_from_directory(DB_PATH.parent,DB_PATH.name,as_attachment=True)

if __name__ == "__main__":
    app.run(host="0.0.0.0",port=5000,debug=False)
